import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
import os



# Step 1: Read the source Excel file
source_file = "./StudentRepositoryData.xlsx"  # Replace with your source file name
sheet_name = "Repository Data"  # Replace with your sheet name if necessary

# Load the source data into a pandas DataFrame
df = pd.read_excel(source_file, sheet_name=sheet_name)

# Step 2: Extract relevant columns (adjust column names as per your data)
required_columns = ["Member Name", "Repository URL", "Java File Count"]
processed_df = df[required_columns]

# Step 3: Create a new Excel file and add the data
# Specify the file path
file_path = "students_table_with_visuals.xlsx"

# Check if the file exists before deleting
if os.path.exists(file_path):
    os.remove(file_path)
    print(f"{file_path} has been deleted.")
else:
    print(f"{file_path} does not exist.")
new_file = "students_table_with_visuals.xlsx"  # Name of the new file

# Using openpyxl to create a workbook
wb = Workbook()

# --- Create Students Table Sheet ---
ws_table = wb.active
ws_table.title = "Students Table"

# Write headers to the sheet
headers = required_columns
ws_table.append(headers)

# Write rows to the sheet
for row in processed_df.itertuples(index=False, name=None):
    ws_table.append(row)

# --- Create Visualization Sheet ---
ws_visual = wb.create_sheet(title="Visualization")

# Add a chart for .java File Count
chart = BarChart()  # You can also use PieChart()

# Define the data for the chart
data = Reference(ws_table, min_col=3, min_row=1, max_row=processed_df.shape[0] + 1, max_col=3)  # "Java File Count" column
categories = Reference(ws_table, min_col=1, min_row=2, max_row=processed_df.shape[0] + 1)  # "Name" column
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)

# Chart title and styling
chart.title = "Java File Count Per Student"
chart.style = 10  # Predefined styles

# Add chart to the Visualization sheet
ws_visual.add_chart(chart, "A1")

# Step 4: Save the file
wb.save(new_file)

print(f"New Excel file with visualization created: {new_file}")
